"""PDF (reportlab), DOCX (python-docx) and XLSX (openpyxl, with native charts) renderers for a report Doc.

Fonts: PDF uses DejaVu Sans (Latin, Greek, Cyrillic, accented text). reportlab does not shape complex scripts: Devanagari and Arabic
text in a PDF is NOT laid out correctly - use DOCX or XLSX for Hindi reports (they store the text as Unicode and let Word/Excel shape it).
Set MKA_REPORT_FONT to a TTF path to use another font for PDF."""
from __future__ import annotations

import io
import os
from functools import lru_cache
from xml.sax.saxutils import escape

from ..analytics import charts
from .model import Doc

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME = {"pdf": PDF_MIME, "docx": DOCX_MIME, "xlsx": XLSX_MIME}
MAX_CELL = 32000


def _s(v) -> str:
    if isinstance(v, float):
        return f"{v:,.2f}".rstrip("0").rstrip(".") if not v.is_integer() else f"{int(v):,}"
    return str(v)


# ----------------------------------------------------------------------------- PDF
@lru_cache(maxsize=4)
def _pdf_fonts(path: str) -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if not path:
        import matplotlib

        base = os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf")
        regular, bold = os.path.join(base, "DejaVuSans.ttf"), os.path.join(base, "DejaVuSans-Bold.ttf")
    else:
        regular = bold = path
    pdfmetrics.registerFont(TTFont("MKA", regular))
    pdfmetrics.registerFont(TTFont("MKA-Bold", bold))
    pdfmetrics.registerFontFamily("MKA", normal="MKA", bold="MKA-Bold", italic="MKA", boldItalic="MKA-Bold")
    return "MKA", "MKA-Bold"


def to_pdf(doc: Doc, font_path: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    reg, bold = _pdf_fonts(font_path)
    base = ParagraphStyle("b", fontName=reg, fontSize=9.5, leading=13)
    styles = {1: ParagraphStyle("h1", parent=base, fontName=bold, fontSize=15, leading=19, spaceBefore=10, spaceAfter=5),
              2: ParagraphStyle("h2", parent=base, fontName=bold, fontSize=12, leading=16, spaceBefore=8, spaceAfter=3),
              3: ParagraphStyle("h3", parent=base, fontName=bold, fontSize=10.5, leading=14, spaceBefore=6, spaceAfter=2)}
    note = ParagraphStyle("n", parent=base, textColor=colors.HexColor("#555555"), fontSize=8.5, leading=11)
    cell = ParagraphStyle("c", parent=base, fontSize=8, leading=10)
    buf = io.BytesIO()

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont(reg, 7.5)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(18 * mm, 10 * mm, f"{doc.title[:80]} — generated {doc.timestamp}")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {d.page}")
        canvas.restoreState()

    out = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=18 * mm, title=doc.title, author=doc.generated_by or "Multi-Agent Knowledge Assistant")
    width = A4[0] - 36 * mm

    def para(text: str, st=base):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), st)

    story = [Paragraph(escape(doc.title), ParagraphStyle("t", parent=base, fontName=bold, fontSize=20, leading=24, spaceAfter=4))]
    if doc.subtitle:
        story.append(para(doc.subtitle, note))
    story += [para(f"Generated {doc.timestamp}" + (f" by {doc.generated_by}" if doc.generated_by else ""), note), Spacer(1, 6)]
    for kind, d in doc.blocks:
        if kind == "heading":
            story.append(para(d["text"], styles[d["level"]]))
        elif kind == "paragraph":
            story += [para(d["text"]), Spacer(1, 4)]
        elif kind == "note":
            story += [para(d["text"], note), Spacer(1, 4)]
        elif kind == "bullets":
            story.append(ListFlowable([ListItem(para(i), leftIndent=10) for i in d["items"]], bulletType="bullet", start="•", leftIndent=12))
            story.append(Spacer(1, 4))
        elif kind == "table":
            hdr, rows = d["header"], d["rows"]
            data = [[Paragraph(f"<b>{escape(h)}</b>", cell) for h in hdr]] + [[Paragraph(escape(_s(c)[:600]), cell) for c in r] for r in rows[:500]]
            t = Table(data, colWidths=[width / len(hdr)] * len(hdr), repeatRows=1)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF7")), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B5BFCC")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
            if d.get("caption"):
                story.append(para(d["caption"], note))
            story += [t, Spacer(1, 6)]
            if len(rows) > 500:
                story.append(para(f"… {len(rows) - 500} more rows not shown (use the XLSX export for the full table).", note))
        elif kind == "chart":
            story += [Image(io.BytesIO(charts.render_png(d["spec"])), width=width, height=width * 0.5), Spacer(1, 6)]
    if doc.sources:
        story.append(para("Sources", styles[1]))
        story.append(_pdf_table(["Document", "Version", "Type", "Indexed", "Hash"], [[s["name"], s["version"], s["kind"], s["ingested_at"], s["hash"]] for s in doc.sources], width, cell, bold))
    out.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _pdf_table(header, rows, width, cell, bold):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    data = [[Paragraph(f"<b>{escape(h)}</b>", cell) for h in header]] + [[Paragraph(escape(_s(c)), cell) for c in r] for r in rows]
    t = Table(data, colWidths=[width * w for w in (0.38, 0.1, 0.1, 0.25, 0.17)], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF7")), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B5BFCC")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


# ----------------------------------------------------------------------------- DOCX
def to_docx(doc: Doc) -> bytes:
    import docx
    from docx.shared import Inches, Pt

    d = docx.Document()
    d.core_properties.title, d.core_properties.author = doc.title, doc.generated_by or "Multi-Agent Knowledge Assistant"
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading(doc.title, 0)
    if doc.subtitle:
        d.add_paragraph(doc.subtitle).italic = True
    d.add_paragraph(f"Generated {doc.timestamp}" + (f" by {doc.generated_by}" if doc.generated_by else ""))
    for kind, b in doc.blocks:
        if kind == "heading":
            d.add_heading(b["text"], b["level"])
        elif kind in ("paragraph", "note"):
            para = d.add_paragraph(b["text"])
            if kind == "note":
                for r in para.runs:
                    r.italic = True
        elif kind == "bullets":
            for i in b["items"]:
                d.add_paragraph(i, style="List Bullet")
        elif kind == "table":
            if b.get("caption"):
                d.add_paragraph(b["caption"]).runs[0].italic = True
            t = d.add_table(rows=1, cols=len(b["header"]))
            t.style = "Light Grid Accent 1"
            for i, h in enumerate(b["header"]):
                t.rows[0].cells[i].text = h
            for r in b["rows"][:2000]:
                cells = t.add_row().cells
                for i, c in enumerate(r):
                    cells[i].text = _s(c)[:MAX_CELL]
        elif kind == "chart":
            d.add_picture(io.BytesIO(charts.render_png(b["spec"])), width=Inches(6.2))
    if doc.sources:
        d.add_heading("Sources", 1)
        t = d.add_table(rows=1, cols=5)
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(("Document", "Version", "Type", "Indexed", "Hash")):
            t.rows[0].cells[i].text = h
        for s in doc.sources:
            cells = t.add_row().cells
            for i, v in enumerate((s["name"], s["version"], s["kind"], s["ingested_at"], s["hash"])):
                cells[i].text = str(v)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------- XLSX
def _safe_cell(v):
    """Strings beginning with = + - @ are written as text (never as formulas): report content comes from documents and must not execute in Excel."""
    if isinstance(v, str):
        v = v[:MAX_CELL]
        return "'" + v if v[:1] in ("=", "+", "-", "@") else v
    return v


def to_xlsx(doc: Doc) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"], ws["A1"].font = doc.title, Font(bold=True, size=16)
    ws["A2"] = f"Generated {doc.timestamp}" + (f" by {doc.generated_by}" if doc.generated_by else "")
    row, fill = 4, PatternFill("solid", fgColor="E8EEF7")
    ws.column_dimensions["A"].width = 38
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 22
    for kind, b in doc.blocks:
        if kind == "heading":
            ws.cell(row, 1, b["text"]).font = Font(bold=True, size=14 - b["level"] * 1.5)
            row += 1
        elif kind in ("paragraph", "note"):
            c = ws.cell(row, 1, _safe_cell(b["text"]))
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row].height = max(15, 15 * (len(b["text"]) // 130 + 1))
            row += 1
        elif kind == "bullets":
            for i in b["items"]:
                ws.cell(row, 1, _safe_cell("• " + i)).alignment = Alignment(wrap_text=True)
                row += 1
        elif kind in ("table", "chart"):
            if kind == "chart":
                sp = b["spec"]
                header, rows = ["label"] + [s["name"] for s in sp["series"]], [[lab] + [s["values"][i] for s in sp["series"]] for i, lab in enumerate(sp["labels"])]
            else:
                header, rows = b["header"], b["rows"][:50000]
                if b.get("caption"):
                    ws.cell(row, 1, b["caption"]).font = Font(italic=True)
                    row += 1
            top = row
            for j, h in enumerate(header, 1):
                c = ws.cell(row, j, _safe_cell(h))
                c.font, c.fill = Font(bold=True), fill
            for r in rows:
                row += 1
                for j, v in enumerate(r, 1):
                    ws.cell(row, j, _safe_cell(v))
            if kind == "chart" and rows:
                sp = b["spec"]
                ch = BarChart() if sp["type"] == "bar" else LineChart()
                ch.title, ch.height, ch.width = sp["title"], 8, 16
                ch.y_axis.title, ch.x_axis.title = sp.get("y_label") or None, sp.get("x_label") or None
                ch.add_data(Reference(ws, min_col=2, max_col=len(header), min_row=top, max_row=row), titles_from_data=True)
                ch.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=row))
                ws.add_chart(ch, f"{chr(ord('A') + min(len(header), 5) + 1)}{top}")
                row = max(row, top + 17)
            row += 2
    if doc.sources:
        s = wb.create_sheet("Sources")
        for j, h in enumerate(("Document", "Version", "Type", "Indexed", "Hash"), 1):
            c = s.cell(1, j, h)
            c.font, c.fill = Font(bold=True), fill
        for i, src in enumerate(doc.sources, 2):
            for j, v in enumerate((src["name"], src["version"], src["kind"], src["ingested_at"], src["hash"]), 1):
                s.cell(i, j, _safe_cell(v))
        s.column_dimensions["A"].width = 44
        s["G1"], s["G2"] = "Generated", doc.timestamp
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render(doc: Doc, fmt: str, font_path: str = "") -> bytes:
    if fmt == "pdf":
        return to_pdf(doc, font_path)
    if fmt == "docx":
        return to_docx(doc)
    if fmt == "xlsx":
        return to_xlsx(doc)
    raise ValueError("format must be pdf, docx or xlsx")
