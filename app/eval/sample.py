"""Synthetic evaluation corpus + question set. Ground truth is computed from the generating data, so it cannot drift.
It is deliberately small and mixed (PDF, DOCX, XLSX, CSV, MD, TXT, JSON): it measures regressions and strategy differences, it is
NOT a benchmark of real-world accuracy."""
from __future__ import annotations

import io
import json

from .dataset import EvalCase, EvalDataset, RelevantRef

HR_PAGES = [
    "HR Policy Handbook. Section 1: Leave. Full-time employees accrue 25 days of annual leave per calendar year. "
    "Up to 5 unused days may be carried over into the next year. Sick leave is limited to 10 days per year and requires a doctor's note after 3 consecutive days.",
    "Section 2: Remote work. Employees may work remotely up to 3 days per week with written manager approval. "
    "Section 3: Expenses. Expense reports must be filed within 30 days of purchase. Receipts are mandatory for any expense above $25. "
    "Travel must be booked through the approved agency.",
    "Section 4: Code of conduct. Harassment of any kind is prohibited and must be reported to HR. Questions about this handbook can be sent to Anjali Rao, the HR Director, "
    "at anjali.rao@acme.example. Anjali Rao also chairs the annual policy review.",
]
SEC_PAGES = [
    "Security Policy. Passwords must be rotated every 90 days and must contain at least 14 characters. Multi-factor authentication is required for all remote access.",
    "Incident response: any suspected security incident must be reported to the security team within 24 hours. Customer data must be retained for 7 years and then securely destroyed.",
]
CONTRACT_PARAS = [
    ("Clause 1. Definitions", "In this Agreement, Services means the software support services described in Schedule A."),
    ("Clause 4. Payment", "The Customer shall pay each invoice within 45 days of the invoice date. Late payments accrue interest at 2 percent per month."),
    ("Clause 8. Termination", "Either party may terminate this Agreement for convenience by giving ninety days written notice to the other party."),
    ("Clause 12. Confidentiality", "Each party shall keep the other party's confidential information secret for a period of 5 years after termination."),
    ("Clause 15. Governing law", "This Agreement is governed by the laws of the State of Delaware."),
]
FEES = [["Plan", "Monthly fee", "Seats"], ["Basic", "$1,000", "10"], ["Pro", "$2,500", "50"], ["Enterprise", "$9,000", "500"]]
SALES = [["Employee", "Region", "Product", "Sales"], ["Alice", "North", "Widget", 12000], ["Bob", "South", "Gadget", 34000], ["Carol", "North", "Gadget", 9000], ["Dan", "East", "Widget", 15500],
         ["Erin", "South", "Widget", 22000], ["Frank", "East", "Gadget", 8000], ["Grace", "North", "Widget", 31000], ["Heidi", "West", "Gadget", 27000], ["Ivan", "West", "Widget", 5000]]
CUSTOMERS = [["ID", "Name", "Email", "Country", "Plan"], [1, "Asha Verma", "asha@mail.example", "India", "Pro"], [2, "Liam Walker", "liam@mail.example", "UK", "Basic"],
             [3, "Mei Chen", "mei@mail.example", "China", "Pro"], [4, "Omar Haddad", "omar@mail.example", "UAE", "Enterprise"], [5, "Asha V.", "ASHA@mail.example", "India", "Basic"],
             [6, "Sofia Rossi", "sofia@mail.example", "Italy", "Pro"], [7, "Noah Smith", "noah@mail.example", "UK", "Pro"]]
NOTES_MD = """# Planning meeting March 2024

Attendees: Anjali Rao, Bob Mehta and Carol Singh from Acme Corp.

## Decisions

The team decided to launch the mobile app in June 2024. Bob Mehta will own the launch checklist. The budget was approved at $50,000 by Anjali Rao.

## Partnerships

Carol Singh signed the reseller agreement with Delta Systems. Delta Systems will distribute the Pro plan in Europe.

## Risks

The main risk is the shortage of qualified testers. Dev Patel from Delta Systems offered two engineers.
"""
FAQ_TXT = """Frequently asked questions

How do I reset my password? Click "Forgot password" on the login page and follow the emailed link. The link expires after 1 hour.

What is the return policy? Products can be returned within 30 days of delivery for a full refund, provided they are unused.

How long does shipping take? Standard shipping takes 5 to 7 business days; express shipping takes 2 business days.

How long is the warranty? All hardware products carry a 24 month limited warranty.
"""
INVENTORY = {"warehouse": "Pune", "items": [{"sku": "SKU-1001", "name": "Steel bolt", "qty": 500}, {"sku": "SKU-1002", "name": "Copper wire", "qty": 120},
                                           {"sku": "SKU-1003", "name": "Rubber gasket", "qty": 0}], "manager": "Dev Patel"}


def _count(term: str) -> int:
    """Whole-word, case-insensitive occurrences of `term` in all generated text (ground truth for exact-search cases)."""
    import re

    text = "\n".join(HR_PAGES + SEC_PAGES + [f"{h}\n{b}" for h, b in CONTRACT_PARAS] + [NOTES_MD, FAQ_TXT, json.dumps(INVENTORY)] + [" ".join(map(str, r)) for r in FEES + SALES + CUSTOMERS])
    return len(re.findall(r"(?<!\w)" + r"\s+".join(map(re.escape, term.split())) + r"(?!\w)(?!\.\d)", text, re.I))


def _pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)  # reproducible bytes: re-loading the demo corpus is detected as a duplicate
    for text in pages:
        y, line = 800, ""
        words = text.split()
        lines = []
        for w in words:
            if len(line) + len(w) > 95:
                lines.append(line)
                line = ""
            line += (" " if line else "") + w
        lines.append(line)
        for ln in lines:
            c.drawString(40, y, ln)
            y -= 14
        c.showPage()
    c.save()
    return buf.getvalue()


def _docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Master Services Agreement", 1)
    for h, body in CONTRACT_PARAS:
        d.add_heading(h, 2)
        d.add_paragraph(body)
    t = d.add_table(rows=len(FEES), cols=3)
    for i, row in enumerate(FEES):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    from datetime import datetime

    wb = Workbook()
    wb.properties.created = wb.properties.modified = datetime(2024, 1, 1)  # reproducible bytes
    ws = wb.active
    ws.title = "Q1"
    for r in SALES:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(rows) -> bytes:
    return ("\n".join(",".join(str(c) for c in r) for r in rows) + "\n").encode()


def corpus() -> dict[str, bytes]:
    return {"hr_policy.pdf": _pdf(HR_PAGES), "security_policy.pdf": _pdf(SEC_PAGES), "service_agreement.docx": _docx(), "sales_q1.xlsx": _xlsx(),
            "customers.csv": _csv(CUSTOMERS), "meeting_notes.md": NOTES_MD.encode(), "product_faq.txt": FAQ_TXT.encode(), "inventory.json": json.dumps(INVENTORY).encode()}


def dataset() -> EvalDataset:
    R = RelevantRef
    sales = [r[3] for r in SALES[1:]]
    north = sum(r[3] for r in SALES[1:] if r[1] == "North")
    cases = [
        # --- semantic: lexical overlap with the source, from easy to hard
        EvalCase(id="s01", expected_route="semantic", question="How many days of annual leave do full-time employees accrue?", relevant=[R(source="hr_policy.pdf", page=1, contains="25 days of annual leave")], expected_keywords=["25 days"]),
        EvalCase(id="s02", expected_route="semantic", question="How many unused leave days can be carried over?", relevant=[R(source="hr_policy.pdf", page=1, contains="carried over")], expected_keywords=["5"]),
        EvalCase(id="s03", expected_route="semantic", question="When is a doctor's note required for sick leave?", relevant=[R(source="hr_policy.pdf", page=1, contains="doctor")], expected_keywords=["3 consecutive days"]),
        EvalCase(id="s04", expected_route="semantic", question="How many days a week can employees work remotely?", relevant=[R(source="hr_policy.pdf", page=2, contains="remotely")], expected_keywords=["3 days"]),
        EvalCase(id="s05", expected_route="semantic", question="Within how many days must expense reports be filed?", relevant=[R(source="hr_policy.pdf", page=2, contains="30 days")], expected_keywords=["30 days"]),
        EvalCase(id="s06", expected_route="semantic", question="Above what amount are receipts mandatory?", relevant=[R(source="hr_policy.pdf", page=2, contains="$25")], expected_keywords=["$25"]),
        EvalCase(id="s07", expected_route="semantic", question="How often must passwords be rotated?", relevant=[R(source="security_policy.pdf", page=1, contains="90 days")], expected_keywords=["90 days"]),
        EvalCase(id="s08", expected_route="semantic", question="Within how many hours must a security incident be reported?", relevant=[R(source="security_policy.pdf", page=2, contains="24 hours")], expected_keywords=["24 hours"]),
        EvalCase(id="s09", expected_route="semantic", question="How long must customer data be retained?", relevant=[R(source="security_policy.pdf", page=2, contains="7 years")], expected_keywords=["7 years"]),
        EvalCase(id="s10", expected_route="semantic", question="What notice is required to terminate the agreement for convenience?", relevant=[R(source="service_agreement.docx", contains="ninety days written notice")], expected_keywords=["ninety days"]),
        EvalCase(id="s11", expected_route="semantic", question="Within how many days must invoices be paid?", relevant=[R(source="service_agreement.docx", contains="45 days")], expected_keywords=["45 days"]),
        EvalCase(id="s12", expected_route="semantic", question="Which law governs the agreement?", relevant=[R(source="service_agreement.docx", contains="Delaware")], expected_keywords=["Delaware"]),
        EvalCase(id="s13", expected_route="semantic", question="How long does the confidentiality obligation last after termination?", relevant=[R(source="service_agreement.docx", contains="5 years")], expected_keywords=["5 years"]),
        EvalCase(id="s14", expected_route="semantic", question="What is the return policy for products?", relevant=[R(source="product_faq.txt", contains="30 days of delivery")], expected_keywords=["30 days"]),
        EvalCase(id="s15", expected_route="semantic", question="How long is the hardware warranty?", relevant=[R(source="product_faq.txt", contains="24 month")], expected_keywords=["24 month"]),
        EvalCase(id="s16", expected_route="semantic", question="When will the mobile app launch?", relevant=[R(source="meeting_notes.md", contains="June 2024")], expected_keywords=["June 2024"]),
        EvalCase(id="s17", expected_route="semantic", question="Who signed the reseller agreement with Delta Systems?", relevant=[R(source="meeting_notes.md", contains="reseller agreement")], expected_keywords=["Carol Singh"]),
        EvalCase(id="s18", expected_route="semantic", question="What is the main risk to the launch?", relevant=[R(source="meeting_notes.md", contains="shortage of qualified testers")], expected_keywords=["testers"]),
        # hard paraphrases (little word overlap) - an offline lexical embedder is expected to miss some of these
        EvalCase(id="h01", expected_route="semantic", question="How much vacation time do staff get yearly?", relevant=[R(source="hr_policy.pdf", page=1, contains="25 days of annual leave")], expected_keywords=["25"]),
        EvalCase(id="h02", expected_route="semantic", question="How soon do we have to tell the security team about a breach?", relevant=[R(source="security_policy.pdf", page=2, contains="24 hours")], expected_keywords=["24"]),
        EvalCase(id="h03", expected_route="semantic", question="What is the cost of the Pro subscription per month?", relevant=[R(source="service_agreement.docx", contains="$2,500")], expected_keywords=["$2,500"]),
        # --- exact
        EvalCase(id="e01", question="Does Anjali Rao appear in the documents?", type="exact", expected_route="exact", expected_count=_count("Anjali Rao"), relevant=[R(source="hr_policy.pdf", page=3, contains="Anjali Rao")]),
        EvalCase(id="e02", question='How many times does "Delta Systems" appear?', type="exact", expected_route="exact", expected_count=_count("Delta Systems")),
        EvalCase(id="e03", question="Find SKU-1003", type="exact", expected_route="exact", expected_count=_count("SKU-1003")),
        EvalCase(id="e04", question='Search for "Zaphod Beeblebrox"', type="exact", expected_route="exact", expected_count=0),
        EvalCase(id="e05", question="Where does clause 8 appear?", type="exact", expected_route="document_search", expected_count=_count("clause 8")),
        EvalCase(id="e06", question='How many times does "Bob Mehta" appear?', type="exact", expected_route="exact", expected_count=_count("Bob Mehta")),
        # --- data
        EvalCase(id="d01", question="What is the total sales?", type="data", expected_route="data", expected_value=float(sum(sales))),
        EvalCase(id="d02", question="Which employee has the highest sales?", type="data", expected_route="data", expected_keywords=["Bob"], expected_value=34000.0),
        EvalCase(id="d03", question="What is the average sales?", type="data", expected_route="data", expected_value=sum(sales) / len(sales)),
        EvalCase(id="d04", question="What is the total sales in the North region?", type="data", expected_route="data", expected_value=float(north)),
        EvalCase(id="d05", question="Which region has the highest total sales?", type="data", expected_route="data", expected_keywords=["South"]),
        EvalCase(id="d06", question="How many customers are in the UK?", type="data", expected_route="data", expected_value=2.0),
        EvalCase(id="d07", question="Are there duplicate emails in customers.csv?", type="data", expected_route="data", expected_value=1.0),
        EvalCase(id="d08", question="Show the top 3 employees sorted by sales", type="data", expected_route="data", expected_keywords=["Bob", "Grace"]),
        EvalCase(id="d09", question="What is the median sales?", type="data", expected_route="data", expected_value=15500.0),
        # --- unanswerable -> must abstain
        EvalCase(id="a01", question="What is the capital of Mongolia?", type="abstain", answerable=False),
        EvalCase(id="a02", question="What was the weather in Pune yesterday?", type="abstain", answerable=False),
        EvalCase(id="a03", question="Who won the football world cup in 1998?", type="abstain", answerable=False),
        EvalCase(id="a04", question="What is the CEO's favourite colour?", type="abstain", answerable=False),
        # --- document intelligence
        EvalCase(id="i01", question="How is Anjali Rao related to Dev Patel?", type="intel", expected_route="summarize", expected_keywords=["Delta Systems"]),
        EvalCase(id="i02", question="Extract all the entities from meeting_notes.md", type="intel", expected_route="summarize", expected_keywords=["Anjali Rao"]),
    ]
    return EvalDataset(name="synthetic-mixed-v1", cases=cases)
