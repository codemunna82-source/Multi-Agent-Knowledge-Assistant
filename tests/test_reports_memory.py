import io
from datetime import UTC, datetime, timedelta

import docx
import openpyxl
import pymupdf
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app import maintenance
from app.api import create_app
from app.config import Settings
from app.db import Conversation, MemoryFact, Message, Report
from app.security.pii import findings, redact
from app.service import Knowledge

from .conftest import SALES_ROWS, make_docx, make_pdf, make_xlsx
from .fakes import FakeLLM

PW = "correct horse battery"
V1 = b"Employees accrue 25 days of annual leave per year. Expense reports must be filed within 30 days.\n\nManagers approve all schedules."
V2 = b"Employees accrue 30 days of annual leave per year. Expense reports must be filed within 30 days.\n\nManagers approve all schedules. A wellness stipend of 500 dollars is offered."


class Env:
    def __init__(self, tmp_path, llm=None, **kw):
        self.s = Settings(data_dir=str(tmp_path / "d"), auth_enabled=True, jwt_secret="x" * 40, rate_limit_ask="1000/minute", rate_limit_auth="1000/minute", rate_limit_ingest="1000/minute", _env_file=None, **kw)
        self.k = Knowledge(self.s, llm=llm)
        self.c = TestClient(create_app(self.s, self.k))

    def tenant(self, name, email, ws="w"):
        tok = self.c.post("/auth/signup", json={"tenant_name": name, "email": email, "password": PW}).json()["access_token"]
        return {"Authorization": f"Bearer {tok}", "X-Session-ID": ws}

    def user(self, admin_h, email, role="user", ws="w"):
        self.c.post("/admin/users", json={"email": email, "password": PW, "role": role}, headers=admin_h)
        tok = self.c.post("/auth/login", json={"email": email, "password": PW}).json()["access_token"]
        return {"Authorization": f"Bearer {tok}", "X-Session-ID": ws}

    def up(self, h, name, data, vis="tenant"):
        r = self.c.post("/sources/upload", files=[("files", (name, data))], data={"visibility": vis}, headers=h)
        assert r.status_code == 200, r.text
        return r.json()[0]["source"]

    def ask(self, h, q, **kw):
        r = self.c.post("/ask", json={"question": q, **kw}, headers=h)
        assert r.status_code == 200, r.text
        return r.json()


@pytest.fixture
def env(tmp_path):
    return Env(tmp_path)


# ----------------------------------------------------------------------------- reports
def pdf_text(data: bytes) -> str:
    return "\n".join(p.get_text() for p in pymupdf.open(stream=data, filetype="pdf"))


def make_report(env, h, **body):
    r = env.c.post("/reports", json=body, headers=h)
    assert r.status_code == 201, r.text
    return r.json()


def test_summary_report_pdf_has_sources_and_timestamps(env):
    h = env.tenant("Acme", "a@acme.io")
    src = env.up(h, "policy.pdf", make_pdf(["Employees accrue 25 days of annual leave per year.", "Expense reports must be filed within 30 days of the trip."]))
    rep = make_report(env, h, kind="summary", format="pdf", source_ids=[src["id"]], params={"style": "detailed"})
    assert rep["kind"] == "summary" and rep["format"] == "pdf" and rep["size_bytes"] > 1500
    dl = env.c.get(rep["download_url"], headers=h)
    assert dl.status_code == 200 and dl.headers["content-type"] == "application/pdf" and dl.content[:5] == b"%PDF-" and dl.headers["content-disposition"].endswith('Summary_report.pdf"')
    txt = pdf_text(dl.content)
    assert "Summary report" in txt and "annual leave" in txt and "Sources" in txt and "policy.pdf" in txt
    assert datetime.now(UTC).strftime("%Y-%m-%d") in txt and "a@acme.io" in txt and "generated" in txt.lower()


def test_comparison_report_docx_lists_changes(env):
    h = env.tenant("Acme", "a@acme.io")
    a, b = env.up(h, "v1.txt", V1), env.up(h, "v2.txt", V2)
    rep = make_report(env, h, kind="comparison", format="docx", source_ids=[a["id"], b["id"]], params={"mode": "diff"})
    d = docx.Document(io.BytesIO(env.c.get(rep["download_url"], headers=h).content))
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert "25→30" in cells and any("wellness stipend" in c for c in cells) and any(c == "v1.txt" for c in cells)
    assert d.paragraphs[0].text == "Comparison report" and any("Generated" in p.text for p in d.paragraphs)
    contra = make_report(env, h, kind="comparison", format="docx", source_ids=[a["id"], b["id"]], params={"mode": "contradictions"})
    cd = docx.Document(io.BytesIO(env.c.get(contra["download_url"], headers=h).content))
    assert any("annual leave" in c.text for t in cd.tables for r in t.rows for c in r.cells)


def test_analytics_report_xlsx_has_data_and_native_chart(env):
    h = env.tenant("Acme", "a@acme.io")
    rows = "Date,Region,Revenue\n" + "\n".join(f"{2023 + i // 12}-{i % 12 + 1:02d}-15,N,{1000 + 10 * i if i != 17 else 9000}" for i in range(24))
    src = env.up(h, "orders.csv", rows.encode())
    rep = make_report(env, h, kind="analytics", format="xlsx", source_ids=[src["id"]], params={"queries": ["SELECT region, SUM(revenue) AS total FROM orders GROUP BY region"]})
    wb = openpyxl.load_workbook(io.BytesIO(env.c.get(rep["download_url"], headers=h).content))
    assert wb.sheetnames == ["Report", "Sources"]
    ws = wb["Report"]
    flat = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Data analysis report" in flat[0] and any("Outliers in Revenue" in x for x in flat) and "2024-06" in flat and any("SUM(revenue)" in x for x in flat)
    assert len(ws._charts) == 1 and ws._charts[0].title is not None
    assert wb["Sources"]["A2"].value == "orders.csv" and wb["Sources"]["H2"].value is None and wb["Sources"]["G2"].value.endswith("UTC")


def test_xlsx_never_writes_formulas_from_document_text(env):
    h = env.tenant("Acme", "a@acme.io")
    src = env.up(h, "evil.txt", b"=HYPERLINK(\"http://evil\",\"click\") is a formula. +cmd|' /C calc'!A0 should stay text.")
    rep = make_report(env, h, kind="summary", format="xlsx", source_ids=[src["id"]], params={"style": "detailed"})
    ws = openpyxl.load_workbook(io.BytesIO(env.c.get(rep["download_url"], headers=h).content))["Report"]
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)


def test_exact_search_and_research_reports(tmp_path):
    env = Env(tmp_path)
    h = env.tenant("Acme", "a@acme.io")
    src = env.up(h, "policy.pdf", make_pdf(["Contact Anjali Rao for help.", "Anjali Rao approves travel.", "Nothing here."]))
    rep = make_report(env, h, kind="exact_search", format="pdf", source_ids=[src["id"]], params={"term": "Anjali Rao"})
    txt = pdf_text(env.c.get(rep["download_url"], headers=h).content)
    assert "Exact search report" in txt and "Anjali Rao" in txt and "TOTAL" in txt and "1, 2" in txt
    rr = make_report(env, h, kind="research", format="docx", source_ids=[src["id"]], params={"questions": ["Who approves travel?", "Who should I contact for help?"], "title": "Travel research"})
    d = docx.Document(io.BytesIO(env.c.get(rr["download_url"], headers=h).content))
    body = "\n".join(p.text for p in d.paragraphs)
    assert "Travel research" in body and "1. Who approves travel?" in body and "2. Who should I contact for help?" in body
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert "Evidence from your documents" in body and "policy.pdf" in cells


def test_report_validation_errors(env):
    h = env.tenant("Acme", "a@acme.io")
    a = env.up(h, "a.txt", V1)
    post = lambda **b: env.c.post("/reports", json=b, headers=h)  # noqa: E731
    assert post(kind="summary", source_ids=["0" * 32]).status_code == 422
    assert post(kind="comparison", source_ids=[a["id"]]).status_code == 422
    assert post(kind="exact_search", source_ids=[a["id"]], params={}).status_code == 422
    assert post(kind="research", source_ids=[a["id"]], params={"questions": []}).status_code == 422
    assert post(kind="analytics", source_ids=[a["id"]]).status_code == 422                  # not a spreadsheet
    assert post(kind="nonsense").status_code == 422 and post(kind="summary", format="exe").status_code == 422
    assert env.c.get("/reports/doesnotexist", headers=h).status_code == 404


def test_report_access_control_and_deletion(env):
    admin = env.tenant("Acme", "a@acme.io")
    alice, bob = env.user(admin, "alice@acme.io", "manager"), env.user(admin, "bob@acme.io")
    src = env.up(alice, "a.txt", V1)
    rep = make_report(env, alice, kind="summary", source_ids=[src["id"]])
    assert env.c.get(rep["download_url"], headers=bob).status_code == 404            # same tenant, different user
    assert env.c.get(rep["download_url"], headers=admin).status_code == 200           # admin sees the whole organisation
    other = env.tenant("Beta", "b@beta.io")
    assert env.c.get(rep["download_url"], headers=other).status_code == 404            # other tenant
    assert env.c.get("/reports", headers=bob).json() == [] and len(env.c.get("/reports", headers=alice).json()) == 1
    r = env.c.get("/reports?limit=1", headers=admin)
    assert r.headers["x-total-count"] == "1"
    path = env.s.data_dir
    assert env.c.delete(f"/reports/{rep['id']}", headers=bob).status_code == 404
    assert env.c.delete(f"/reports/{rep['id']}", headers=alice).status_code == 200 and env.c.get(rep["download_url"], headers=alice).status_code == 404
    assert not list((__import__("pathlib").Path(path) / "tenants").rglob("*.pdf"))


def test_reports_respect_document_permissions(env):
    admin = env.tenant("Acme", "a@acme.io")
    alice, bob = env.user(admin, "alice@acme.io", "manager"), env.user(admin, "bob@acme.io", "manager")
    secret = env.up(alice, "secret.txt", V1, vis="private")
    r = env.c.post("/reports", json={"kind": "summary", "source_ids": [secret["id"]]}, headers=bob)
    assert r.status_code == 422 and "inaccessible" in r.json()["detail"]
    ok = make_report(env, bob, kind="summary") if env.up(bob, "mine.txt", V2) else None
    txt = pdf_text(env.c.get(ok["download_url"], headers=bob).content)
    assert "mine.txt" in txt and "secret.txt" not in txt


def test_hindi_report_docx_and_xlsx_keep_unicode_and_pdf_does_not_crash(env):
    h = env.tenant("Acme", "a@acme.io")
    src = env.up(h, "नीति.txt", "कर्मचारियों को साल में पच्चीस दिन की छुट्टी मिलती है। घर से काम करने की अनुमति सप्ताह में तीन दिन है।".encode())
    d = docx.Document(io.BytesIO(env.c.get(make_report(env, h, kind="summary", format="docx", source_ids=[src["id"]])["download_url"], headers=h).content))
    assert any("छुट्टी" in p.text for p in d.paragraphs)
    ws = openpyxl.load_workbook(io.BytesIO(env.c.get(make_report(env, h, kind="summary", format="xlsx", source_ids=[src["id"]])["download_url"], headers=h).content))["Report"]
    assert any("छुट्टी" in str(c.value) for row in ws.iter_rows() for c in row)
    assert env.c.get(make_report(env, h, kind="summary", format="pdf", source_ids=[src["id"]])["download_url"], headers=h).content[:5] == b"%PDF-"


def test_retention_removes_report_files(env):
    h = env.tenant("Acme", "a@acme.io")
    src = env.up(h, "a.txt", V1)
    rep = make_report(env, h, kind="summary", source_ids=[src["id"]])
    env.c.put("/admin/tenant-settings", json={"retention_days_reports": 7, "web_allowed": True, "web_allowed_domains": [], "redact_stored_messages": False}, headers=h)
    with env.k.db.session() as s:
        row = s.get(Report, rep["id"])
        row.created_at = datetime.now(UTC) - timedelta(days=30)
        path = row.path
    assert maintenance.run(env.k, ["retention_reports"], dry_run=True)["retention_reports"]["removed"] == 1
    assert maintenance.run(env.k, ["retention_reports"])["retention_reports"]["removed"] == 1
    assert env.c.get(f"/reports/{rep['id']}", headers=h).status_code == 404 and not __import__("pathlib").Path(path).exists()


# ----------------------------------------------------------------------------- API groups
def test_exact_search_api_counts_pages_and_paginates(env):
    h = env.tenant("Acme", "a@acme.io")
    env.up(h, "policy.pdf", make_pdf([f"Page {i} mentions Anjali Rao here." if i % 2 else f"Page {i} is plain." for i in range(1, 11)]))
    r = env.c.post("/search/exact?limit=2&offset=1", json={"term": "Anjali Rao"}, headers=h)
    j = r.json()
    assert r.status_code == 200 and j["count"] == 5 and r.headers["x-total-count"] == "5" and j["pages"] == {"policy.pdf": [1, 3, 5, 7, 9]} and j["pages_listing_complete"]
    assert [m["loc"]["page"] for m in j["matches"]] == [3, 5] and j["by_source"][0]["count"] == 5
    assert env.c.post("/search/exact", json={"term": "zzzz"}, headers=h).json()["count"] == 0
    assert env.c.post("/search/exact", json={"term": ""}, headers=h).status_code == 422
    assert env.c.post("/search/exact", json={"term": "x", "source_ids": ["0" * 32]}, headers=h).status_code == 404


def test_summaries_api_styles_chapters_and_sections(env):
    h = env.tenant("Acme", "a@acme.io")
    md = (b"# Leave\nEmployees accrue 25 days of annual leave per year. Unused leave expires after twelve months.\n\n# Travel\nTrips must be approved by a manager. "
          b"Expense reports are due within 30 days of returning from a trip.\n\n# Security\nAll laptops must use disk encryption. Passwords rotate every 90 days.\n")
    src = env.up(h, "handbook.md", md)
    post = lambda **b: env.c.post("/summaries", json={"source_ids": [src["id"]], **b}, headers=h)  # noqa: E731
    ch = post(style="chapter").json()["summaries"][0]
    assert [s["title"] for s in ch["sections"]] == ["Leave", "Travel", "Security"] and "25 days" in ch["sections"][0]["summary"] and ch["method"] == "extractive"
    sec = post(style="section", section="travel").json()["summaries"][0]
    assert len(sec["sections"]) == 1 and "manager" in sec["sections"][0]["summary"]
    assert post(style="section", section="nonexistent").status_code == 404 and post(style="section").status_code == 422
    short, detailed, execu = (post(style=s).json()["summaries"][0] for s in ("short", "detailed", "executive"))
    assert len(short["summary"]) <= len(detailed["summary"]) and execu["summary"].startswith("Overview:") and short["style"] == "short"
    assert post(style="poem").status_code == 422


def test_summary_styles_with_a_language_model(tmp_path):
    seen = []
    env = Env(tmp_path, llm=FakeLLM(lambda sy, u: seen.append(u) or "LLM SUMMARY"))
    h = env.tenant("Acme", "a@acme.io")
    src = env.up(h, "a.txt", V1)
    r = env.c.post("/summaries", json={"source_ids": [src["id"]], "style": "executive"}, headers=h).json()
    assert r["llm"] is True and r["summaries"][0]["summary"] == "LLM SUMMARY" and any("Purpose" in u or "executive" in u.lower() for u in seen)


def test_compare_and_changes_api(env):
    h = env.tenant("Acme", "a@acme.io")
    a, b = env.up(h, "a.txt", V1), env.up(h, "b.txt", V2)
    d = env.c.post("/compare", json={"source_ids": [a["id"], b["id"]]}, headers=h).json()
    assert d["summary"]["changed"] == 1 and d["changed"][0]["number_changes"] == [{"from": "25", "to": "30"}] and d["left"]["name"] == "a.txt"
    c = env.c.post("/compare", json={"source_ids": [a["id"], b["id"]], "mode": "contradictions"}, headers=h).json()
    assert c["method"] == "heuristic" and c["contradictions"] and c["contradictions"][0]["type"] == "numeric" and "location" in c["contradictions"][0]["a"]
    assert env.c.post("/compare", json={"source_ids": [a["id"]]}, headers=h).status_code == 422
    v1 = env.up(h, "policy.txt", V1 + b" Version one marker.")
    v2 = env.up(h, "policy.txt", V2 + b" Version two marker.")
    ch = env.c.get(f"/sources/{v2['id']}/changes", headers=h).json()
    assert (ch["from_version"], ch["to_version"], ch["from_id"]) == (1, 2, v1["id"]) and ch["summary"]["changed"] >= 1 and ch["summary"]["added"] >= 1
    assert env.c.get(f"/sources/{v1['id']}/changes", headers=h).status_code == 409
    assert env.c.get(f"/sources/{v2['id']}/changes?against=nope", headers=h).status_code == 404


def test_web_research_endpoint_uses_the_web_agent_only(env):
    h = env.tenant("Acme", "a@acme.io")
    env.up(h, "a.txt", V1)
    r = env.c.post("/web/research", json={"query": "What is the latest news about widgets?"}, headers=h).json()
    assert r["route"] == "web" and "Web search is unavailable" in r["answer"]                    # no search provider configured: honest message, no document answer
    env.c.put("/admin/tenant-settings", json={"retention_days_documents": None, "retention_days_conversations": None, "retention_days_audit": None, "retention_days_reports": None,
                                              "web_allowed": False, "web_allowed_domains": [], "redact_stored_messages": False}, headers=h)
    assert "disabled for your organisation" in env.c.post("/web/research", json={"query": "What does https://example.com/a say?"}, headers=h).json()["answer"]
    assert env.c.get("/conversations", headers=h).json() == []                                  # web research leaves no conversation behind


# ----------------------------------------------------------------------------- memory
def test_pii_helpers():
    assert findings("I prefer short answers") == []
    assert set(findings("mail ann@example.com, card 4111 1111 1111 1111, password: hunter2, 123-45-6789")) >= {"email address", "payment card number", "secret or password", "national ID number"}
    assert "payment card number" not in findings("order number 4111 1111 1111 1112") and findings("order number 1234567") == []   # fails the Luhn check: not a card
    assert redact("ann@example.com") == "[email]" and "4111" not in redact("pay 4111 1111 1111 1111 now") and "hunter2" not in redact("password: hunter2")


def test_memory_notes_crud_validation_and_isolation(env):
    admin = env.tenant("Acme", "a@acme.io")
    alice, bob = env.user(admin, "alice@acme.io"), env.user(admin, "bob@acme.io")
    f = env.c.post("/memory/facts", json={"text": "I prefer answers as short bullet points"}, headers=alice)
    assert f.status_code == 201
    for bad in ("my password is hunter2", "email me at ann@example.com", "card 4111 1111 1111 1111", "x" * 301):
        assert env.c.post("/memory/facts", json={"text": bad}, headers=alice).status_code == 422, bad
    assert [x["text"] for x in env.c.get("/memory/facts", headers=alice).json()] == ["I prefer answers as short bullet points"]
    assert env.c.get("/memory/facts", headers=bob).json() == []
    assert env.c.delete(f"/memory/facts/{f.json()['id']}", headers=bob).status_code == 404
    assert env.c.delete(f"/memory/facts/{f.json()['id']}", headers=alice).status_code == 200
    for i in range(50):
        assert env.c.post("/memory/facts", json={"text": f"note {i}"}, headers=alice).status_code == 201
    assert env.c.post("/memory/facts", json={"text": "one too many"}, headers=alice).status_code == 422
    assert env.c.get("/memory/facts?limit=10", headers=alice).headers["x-total-count"] == "50"
    other = env.tenant("Beta", "b@beta.io")
    assert env.c.get("/memory/facts", headers=other).json() == []


def test_notes_reach_the_model_only_when_long_term_memory_is_on_and_never_as_evidence(tmp_path):
    prompts = []
    env = Env(tmp_path, llm=FakeLLM(lambda sy, u: prompts.append(u) or ("Employees accrue 25 days of annual leave per year [1]." if "Answer the question" in sy else None)))
    h = env.tenant("Acme", "a@acme.io")
    env.up(h, "p.txt", b"Employees accrue 25 days of annual leave per year.")
    env.c.post("/memory/facts", json={"text": "I am a part-time employee working in Pune"}, headers=h)
    q = "How many days of annual leave do employees get?"
    r = env.ask(h, q)
    assert not any("part-time" in p for p in prompts) and "memory_notes_used" not in r["details"]       # off by default
    assert env.c.put("/memory/preferences", json={"long_term_memory": True}, headers=h).json() == {"memory_enabled": True, "long_term_memory": True}
    prompts.clear()
    r = env.ask(h, "How many days of annual leave do employee get as a part-time employee?")
    assert any("part-time employee working in Pune" in p and "NOT evidence" in p for p in prompts) and r["details"]["memory_notes_used"] == 1
    assert all("Pune" not in c["excerpt"] for c in r["citations"])                                   # notes are context, never citations
    prompts.clear()
    r = env.ask(h, "How many days of annual leave do employees get as a part-time employee?", incognito=True)
    assert not any("Pune" in p for p in prompts) and r["details"]["private"] is True


def test_incognito_and_private_conversations_store_nothing(env):
    h = env.tenant("Acme", "a@acme.io")
    env.up(h, "p.txt", V1)
    r = env.ask(h, "How many days of annual leave do employees get?", incognito=True)
    assert r["conversation_id"] is None and env.c.get("/conversations", headers=h).json() == []
    conv = env.c.post("/conversations", headers=h).json()
    assert conv["memory_enabled"] is True
    assert env.c.patch(f"/conversations/{conv['id']}", json={"memory_enabled": False}, headers=h).json()["memory_enabled"] is False
    r = env.ask(h, "How many days of annual leave do employees get?", conversation_id=conv["id"])
    assert r["conversation_id"] is None and env.c.get(f"/conversations/{conv['id']}/messages", headers=h).json() == []
    env.c.patch(f"/conversations/{conv['id']}", json={"memory_enabled": True, "title": "Leave"}, headers=h)
    env.ask(h, "How many days of annual leave do employees get?", conversation_id=conv["id"])
    assert len(env.c.get(f"/conversations/{conv['id']}/messages", headers=h).json()) == 2
    assert env.c.put("/memory/preferences", json={"memory_enabled": False}, headers=h).json()["memory_enabled"] is False
    n = len(env.c.get("/conversations", headers=h).json())
    env.ask(h, "How many days of annual leave do employees get?")
    assert len(env.c.get("/conversations", headers=h).json()) == n and len(env.c.get(f"/conversations/{conv['id']}/messages", headers=h).json()) == 2
    assert env.c.patch("/conversations/zzz", json={"title": "x"}, headers=h).status_code == 404


def test_streaming_respects_incognito(env):
    h = env.tenant("Acme", "a@acme.io")
    env.up(h, "p.txt", V1)
    with env.c.stream("POST", "/ask/stream", json={"question": "How many days of annual leave do employees get?", "incognito": True}, headers=h) as r:
        body = "".join(r.iter_text())
    assert "event: final" in body and env.c.get("/conversations", headers=h).json() == []


def test_redact_stored_messages_setting(env):
    h = env.tenant("Acme", "a@acme.io")
    env.up(h, "p.txt", V1)
    env.c.put("/admin/tenant-settings", json={"retention_days_documents": None, "retention_days_conversations": None, "retention_days_audit": None, "retention_days_reports": None,
                                              "web_allowed": True, "web_allowed_domains": [], "redact_stored_messages": True}, headers=h)
    r = env.ask(h, "My email is ann@example.com and my number is +91 98765 43210; how many days of annual leave do employees get?")
    msgs = env.c.get(f"/conversations/{r['conversation_id']}/messages", headers=h).json()
    assert "ann@example.com" not in msgs[0]["content"] and "[email]" in msgs[0]["content"] and "[phone]" in msgs[0]["content"]


def test_memory_reset_and_user_deletion_erase_personal_data(env):
    admin = env.tenant("Acme", "a@acme.io")
    alice = env.user(admin, "alice@acme.io", "manager")
    src = env.up(alice, "p.txt", V1)
    env.c.post("/memory/facts", json={"text": "I like short answers"}, headers=alice)
    for _ in range(2):
        env.ask(alice, "How many days of annual leave do employees get?")
    rep = make_report(env, alice, kind="summary", source_ids=[src["id"]])
    with env.k.db.session() as s:
        s.get(Conversation, s.scalar(select(Conversation.id))).summary = "old summary"
    r = env.c.delete("/memory", headers=alice).json()
    assert r["facts_deleted"] == 1 and r["conversations_deleted"] == 0 and r["summaries_cleared"] >= 1
    with env.k.db.session() as s:
        assert s.scalars(select(MemoryFact)).all() == [] and all(c.summary is None for c in s.scalars(select(Conversation)))
        assert s.scalars(select(Message)).all()
    assert env.c.delete("/memory?conversations=true", headers=alice).json()["conversations_deleted"] == 1
    with env.k.db.session() as s:
        assert s.scalars(select(Message)).all() == []
    env.c.post("/memory/facts", json={"text": "another note"}, headers=alice)
    uid = env.c.get("/auth/me", headers=alice).json()["id"]
    with env.k.db.session() as s:
        report_file = __import__("pathlib").Path(s.get(Report, rep["id"]).path)
    assert report_file.is_file()
    assert env.c.delete(f"/admin/users/{uid}", headers=admin).status_code == 200
    with env.k.db.session() as s:
        assert s.scalars(select(MemoryFact)).all() == [] and s.get(Report, rep["id"]) is None
    assert not report_file.exists()


def test_memory_preferences_need_a_real_user_when_auth_is_off(client):
    assert client.get("/memory/preferences").json() == {"memory_enabled": True, "long_term_memory": False}
    assert client.put("/memory/preferences", json={"long_term_memory": True}).status_code == 409
    assert client.post("/memory/facts", json={"text": "hello"}).status_code == 409


def test_xls_pdf_sources_are_listed_in_report_sources(env):
    h = env.tenant("Acme", "a@acme.io")
    src = env.up(h, "sales.xlsx", make_xlsx({"Q1": SALES_ROWS}))
    env.up(h, "doc.docx", make_docx(["Hello report world, this is a paragraph of text."]))
    rep = make_report(env, h, kind="summary", format="docx", source_ids=[src["id"]])
    d = docx.Document(io.BytesIO(env.c.get(rep["download_url"], headers=h).content))
    assert any("Spreadsheet with" in p.text for p in d.paragraphs)
