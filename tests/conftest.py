import io
import socket

import pytest
from fastapi.testclient import TestClient

from app.api import create_app
from app.config import Settings


@pytest.fixture
def settings(tmp_path):
    return Settings(data_dir=str(tmp_path / "data"), _env_file=None)


@pytest.fixture
def client(settings):
    return TestClient(create_app(settings))


def make_pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            c.drawString(50, y, line)
            y -= 16
        c.showPage()
    c.save()
    return buf.getvalue()


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_docx(paras: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Policy", 1)
    for p in paras:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, v in enumerate(row):
                t.cell(i, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def upload(client, name, data, session="s1"):
    r = client.post("/sources/upload", files=[("files", (name, data))], headers={"X-Session-ID": session})
    assert r.status_code == 200, r.text
    return r.json()[0]


def ask(client, q, session="s1", **kw):
    r = client.post("/ask", json={"question": q, **kw}, headers={"X-Session-ID": session})
    assert r.status_code == 200, r.text
    return r.json()


SALES_ROWS = [["Employee", "Region", "Sales"], ["Alice", "North", 1200], ["Bob", "South", 3400.5], ["Carol", "North", 900], ["Dave", "South", 3400.5], ["Eve", "East", 150]]
