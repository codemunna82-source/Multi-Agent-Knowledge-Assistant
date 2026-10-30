import json
import math
import sys
import types

import pytest
from fastapi.testclient import TestClient

from app.api import create_app
from app.config import Settings
from app.eval import metrics as M
from app.eval.dataset import EvalCase, EvalDataset, RelevantRef
from app.eval.ragas_adapter import ragas_scores
from app.eval.report import html_report, markdown
from app.eval.runner import compare_strategies, run_eval
from app.eval.__main__ import STRATEGIES, build_sample
from app.security.auth import ANONYMOUS
from app.service import Knowledge

from .conftest import upload
from .fakes import FakeLLM

H = {"X-Session-ID": "s1"}


# ------------------------------------------------------------------ metric math
def test_retrieval_metrics_hand_computed():
    #           rank1   rank2   rank3  rank4    gold passages: 0,1,2
    m = [{0}, set(), {1}, {0}]
    assert M.recall_at_k(m, 1, 3) == pytest.approx(1 / 3)
    assert M.recall_at_k(m, 3, 3) == pytest.approx(2 / 3)
    assert M.recall_at_k(m, 4, 3) == pytest.approx(2 / 3)  # passage 0 found twice counts once
    assert M.precision_at_k(m, 3) == pytest.approx(2 / 3)
    assert M.precision_at_k(m, 4) == pytest.approx(3 / 4)
    assert M.precision_at_k([{0}], 5) == pytest.approx(1 / 5)  # fewer results than k are penalised
    assert M.mrr(m) == 1.0 and M.mrr([set(), set(), {2}]) == pytest.approx(1 / 3) and M.mrr([set()]) == 0.0 and M.mrr([]) == 0.0
    assert M.recall_at_k(m, 3, 0) is None and M.ndcg_at_k(m, 3, 0) is None
    ideal = 1 + 1 / math.log2(3) + 1 / math.log2(4)
    assert M.ndcg_at_k(m, 3, 3) == pytest.approx((1 + 1 / math.log2(4)) / ideal)  # gains at ranks 1 and 3
    assert M.ndcg_at_k([{0}, {1}], 5, 2) == pytest.approx(1.0)


def test_percentiles_and_mean():
    xs = [10, 20, 30, 40, 100]
    assert M.percentile(xs, 50) == 30 and M.percentile(xs, 95) == pytest.approx(88.0) and M.percentile([], 50) is None and M.percentile([7], 95) == 7
    assert M.mean([1, None, 3]) == 2 and M.mean([None]) is None


def test_faithfulness_proxy_detects_unsupported_claims():
    ev = "Employees accrue 25 days of annual leave per year. Unused days carry over up to 5."
    assert M.faithfulness_lexical("Employees accrue 25 days of annual leave per year.", ev) == 1.0
    assert M.faithfulness_lexical("Employees accrue 40 days of annual leave per year.", ev) == 0.0  # wrong number
    assert M.faithfulness_lexical("Employees accrue 25 days of annual leave. The CEO resigned yesterday in Paris.", ev) == 0.5
    assert M.faithfulness_lexical("Most relevant passages found in the indexed sources:\n- Employees accrue 25 days of annual leave per year. [hr.pdf, page 1]", ev) == 1.0
    assert M.faithfulness_lexical("", ev) is None


def test_answer_relevancy_proxy():
    from app.providers import HashEmbedder

    e = HashEmbedder()
    on = M.answer_relevancy_embedding(e, "How many leave days are accrued?", "Employees accrue 25 leave days per year.")
    off = M.answer_relevancy_embedding(e, "How many leave days are accrued?", "The cafeteria serves pasta on Fridays.")
    assert 0 <= off < on <= 1


# ------------------------------------------------------------------ runner on the synthetic corpus (regression guard)
@pytest.fixture(scope="module")
def sample_report(tmp_path_factory):

    td = tmp_path_factory.mktemp("evalsample")
    a, scope, ds = build_sample(Settings(data_dir=str(td / "data"), _env_file=None), td)
    rep = run_eval(a, scope, ANONYMOUS, ds, None)
    rep["comparison"] = compare_strategies(a, scope, ANONYMOUS, ds, STRATEGIES)["comparison"]
    return rep


def test_sample_ground_truth_is_consistent():
    from app.eval.sample import _count, dataset

    ds = dataset()
    ids = [c.id for c in ds.cases]
    assert len(ids) == len(set(ids)) >= 40
    counts = {c.id: c.expected_count for c in ds.cases if c.expected_count is not None}
    assert counts["e01"] == _count("Anjali Rao") == 4 and counts["e02"] == 3 and counts["e04"] == 0


def test_sample_evaluation_regression_thresholds(sample_report):
    m = sample_report["metrics"]
    assert m["cases"] >= 40
    r, a = m["retrieval"], m["answers"]
    assert r["recall@5"] >= 0.85 and r["mrr"] >= 0.85 and r["cases_with_gold"] >= 20
    assert r["recall@5"] >= r["recall@3"] >= r["recall@1"]
    assert a["route_accuracy"] == 1.0 and a["exact_count_accuracy"] == 1.0  # counts come from code, so they must be exact
    assert a["data_accuracy"] >= 0.85 and a["abstained_or_flagged_low_confidence"] == 1.0
    assert a["faithfulness_lexical (proxy)"] >= 0.8
    assert m["latency_ms"]["p50"] is not None and m["latency_ms"]["p95"] >= m["latency_ms"]["p50"]
    assert m["cost_usd"]["total"] == 0 and m["tokens"]["prompt"] == 0  # offline run: no LLM, hash embedder
    assert {"semantic", "exact", "data", "abstain", "intel"} == set(m["by_type"])


def test_sample_failures_are_reported_not_hidden(sample_report):
    ids = {f["id"] for f in sample_report["failures"]}
    assert "h01" in ids  # hard paraphrase the offline lexical embedder cannot bridge
    assert all(f["why"] for f in sample_report["failures"])
    assert len(sample_report["comparison"]) == len(STRATEGIES)


def test_reports_render_and_escape(sample_report):
    md = markdown(sample_report)
    assert "Recall@5" in md and "Strategy comparison" in md and "Failures" in md
    evil = json.loads(json.dumps(sample_report))
    evil["cases"][0]["question"] = '<script>alert(1)</script>'
    evil["name"] = '"><img src=x onerror=alert(1)>'
    h = html_report(evil)
    assert "<script>alert" not in h and "<img src=x" not in h
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in h and "<img src=x" not in h
    assert h.startswith("<!doctype html>") and "Recall@5" in h


def test_usage_cost_and_llm_judge_recorded(tmp_path):
    def script(sy, u):
        if "Answer the question" in sy:
            return "Employees accrue 25 days of annual leave per year [1]."
        if "Rate the answer" in sy:
            return '{"faithfulness": 0.9, "relevancy": 0.8}'
        return None

    s = Settings(data_dir=str(tmp_path / "d"), price_table="fake=1.0:2.0", _env_file=None)
    k = Knowledge(s, llm=FakeLLM(script))
    k.llm.name = "fake:fake"
    from app.agents import Assistant
    from app.storage import Scope

    p = tmp_path / "a.txt"
    p.write_text("Employees accrue 25 days of annual leave per year.")
    scope = Scope("t", "w")
    k.ingest_file(scope, p, "a.txt")
    ds = EvalDataset(name="x", cases=[EvalCase(id="c1", question="How many days of annual leave do employees accrue?", relevant=[RelevantRef(source="a.txt", contains="25 days")], expected_keywords=["25 days"])])
    rep = run_eval(Assistant(k), scope, ANONYMOUS, ds)
    m = rep["metrics"]
    assert m["tokens"]["prompt"] > 0 and m["tokens"]["completion"] > 0 and m["tokens"]["estimated"] is True
    assert m["cost_usd"]["total"] > 0 and m["answers"]["faithfulness_llm_judge"] == 0.9 and m["answers"]["answer_relevancy_llm_judge"] == 0.8
    expected = (m["tokens"]["prompt"] * 1.0 + m["tokens"]["completion"] * 2.0) / 1e6
    assert m["cost_usd"]["total"] == pytest.approx(expected, rel=0.25)  # judge call tokens are tracked separately, so allow slack


def test_ragas_adapter_plumbing_with_stub_module(monkeypatch):
    assert ragas_scores({"cases": []}, {}, {}) is None  # ragas not installed here
    seen = {}
    ds_mod, ragas_mod, metrics_mod = types.ModuleType("datasets"), types.ModuleType("ragas"), types.ModuleType("ragas.metrics")
    ds_mod.Dataset = types.SimpleNamespace(from_dict=lambda d: seen.update(ds=d) or d)
    ragas_mod.evaluate = lambda ds, metrics: {"faithfulness": 0.75, "answer_relevancy": 0.5}
    for n in ("faithfulness", "answer_relevancy", "context_precision", "context_recall"):
        setattr(metrics_mod, n, n)
    monkeypatch.setitem(sys.modules, "datasets", ds_mod)
    monkeypatch.setitem(sys.modules, "ragas", ragas_mod)
    monkeypatch.setitem(sys.modules, "ragas.metrics", metrics_mod)
    rep = {"cases": [{"id": "c1", "question": "q", "answer": "a", "verified": True}, {"id": "c2", "question": "q2", "answer": "a2", "verified": False}]}
    assert ragas_scores(rep, {"c1": ["ctx"], "c2": ["ctx"]}, {"c1": "gt"}) == {"faithfulness": 0.75, "answer_relevancy": 0.5}
    assert seen["ds"] == {"question": ["q"], "answer": ["a"], "contexts": [["ctx"]], "ground_truth": ["gt"]}


# ------------------------------------------------------------------ API
def test_eval_api_sample_run_report_and_cleanup(tmp_path):
    s = Settings(data_dir=str(tmp_path / "d"), auth_enabled=True, jwt_secret="e" * 40, rate_limit_auth="1000/minute", rate_limit_ingest="1000/minute", _env_file=None)
    c = TestClient(create_app(s))
    adm = {"Authorization": "Bearer " + c.post("/auth/signup", json={"tenant_name": "Tenant", "email": "a@a.io", "password": "long enough pw"}).json()["access_token"]}
    c.post("/admin/users", json={"email": "v@a.io", "password": "long enough pw", "role": "viewer"}, headers=adm)
    viewer = {"Authorization": "Bearer " + c.post("/auth/login", json={"email": "v@a.io", "password": "long enough pw"}).json()["access_token"]}
    assert c.post("/eval/run", json={"sample": True}, headers=viewer).status_code == 403
    assert c.post("/eval/run", json={}, headers=adm).status_code == 422
    r = c.post("/eval/run", json={"sample": True, "name": "ci", "compare_strategies": True}, headers=adm)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["metrics"]["retrieval"]["recall@5"] >= 0.85 and len(rep["comparison"]) == 4
    assert c.get("/sources", headers={**adm, "X-Session-ID": "default"}).json() == []
    assert not any((tmp_path / "d" / "tenants").rglob("eval-*"))  # throw-away workspace erased
    runs = c.get("/eval/runs", headers=adm).json()
    assert runs[0]["id"] == rep["id"] and runs[0]["name"] == "ci"
    html = c.get(f"/eval/runs/{rep['id']}/report.html", headers=adm)
    assert html.status_code == 200 and "Recall@5" in html.text and "default-src 'none'" in html.headers["content-security-policy"]
    assert "Recall@5" in c.get(f"/eval/runs/{rep['id']}/report.md", headers=adm).text
    # other tenants cannot read it
    other = {"Authorization": "Bearer " + c.post("/auth/signup", json={"tenant_name": "Other Co", "email": "o@o.io", "password": "long enough pw"}).json()["access_token"]}
    assert c.get(f"/eval/runs/{rep['id']}", headers=other).status_code == 404 and c.get("/eval/runs", headers=other).json() == []


def test_eval_api_custom_dataset_on_own_documents(client):
    upload(client, "policy.txt", b"Employees accrue 25 days of annual leave per year.")
    ds = {"name": "mine", "cases": [{"id": "q1", "question": "How many days of annual leave?", "relevant": [{"source": "policy.txt", "contains": "25 days"}], "expected_keywords": ["25 days"]},
                                     {"id": "q2", "question": "Who won the 1998 world cup?", "type": "abstain", "answerable": False}]}
    r = client.post("/eval/run", json={"dataset": ds}, headers=H)
    assert r.status_code == 200
    m = r.json()["metrics"]
    assert m["retrieval"]["recall@1"] == 1.0 and m["answers"]["keyword_accuracy"] == 1.0 and m["answers"]["abstention_accuracy"] == 1.0
    assert client.post("/eval/run", json={"dataset": {"name": "bad", "cases": []}}, headers=H).status_code == 422


def test_evaluation_is_deterministic_across_hash_seeds(tmp_path):
    """Retrieval must not depend on PYTHONHASHSEED (set/dict ordering): run the sample twice in fresh interpreters with different seeds."""
    import os
    import subprocess

    outs = []
    for seed in ("11", "97"):
        out = tmp_path / f"r{seed}"
        subprocess.run([sys.executable, "-m", "app.eval", "run", "--sample", "--out", str(out)], check=True, capture_output=True, env={**os.environ, "PYTHONHASHSEED": seed})
        rep = json.loads(out.with_suffix(".json").read_text())
        m = {k: v for k, v in rep["metrics"].items() if k != "latency_ms"}
        m["by_type"] = {t: {k: v for k, v in d.items() if k != "latency_p50_ms"} for t, d in m["by_type"].items()}  # wall-clock timings legitimately vary
        outs.append((m, [(c["id"], c["route"], c["verified"], c["answer"], c.get("recall@5"), c.get("mrr")) for c in rep["cases"]]))
    assert outs[0] == outs[1]


# ------------------------------------------------------------------ v3: extended corpus, citation / agent metrics, exports
@pytest.fixture(scope="module")
def extended_report(tmp_path_factory):
    td = tmp_path_factory.mktemp("evalext")
    a, scope, ds = build_sample(Settings(data_dir=str(td / "data"), _env_file=None), td, extended=True)
    return run_eval(a, scope, ANONYMOUS, ds, None)


def test_extended_corpus_ground_truth_and_versions(extended_report):
    from app.eval.extended import _count, corpus, dataset

    names = [n for n, _ in corpus()]
    assert names.count("leave_policy.txt") == 2                                              # version 1 then version 2
    assert _count("Anjali Rao") == 1 and _count("30 days") == 1 and _count("25 days") == 0   # the OLD version is not counted
    assert {c.type for c in dataset().cases} == {"semantic", "compare", "contradiction", "hindi", "analytics", "parallel", "abstain"}
    assert {(c["name"], c["version"]) for c in extended_report["corpus"]} >= {("leave_policy.txt", 2), ("hindi_policy.txt", 1)}


def test_extended_regression_thresholds(extended_report):
    m = extended_report["metrics"]
    a, ag = m["answers"], m["agents"]
    assert a["route_accuracy"] == 1.0 and a["keyword_accuracy"] == 1.0 and a["exact_count_accuracy"] == 1.0 and a["data_accuracy"] == 1.0
    assert a["abstention_accuracy"] == 1.0 and a["false_answer_rate_on_unanswerable"] == 0.0
    assert a["required_sources_cited"] == 1.0 and a["stale_or_forbidden_content_absent"] == 1.0   # the superseded "25 days" never reaches an answer
    assert a["citation_accuracy"] == 1.0 and a["citation_gold_precision"] >= 0.75
    assert ag["execution_success_rate"] == 1.0 and ag["failed_tool_calls"] == 0 and ag["tool_calls"] > 0 and ag["cases_with_invalid_citations"] == 0
    assert m["retrieval"]["recall@5"] >= 0.9 and m["retrieval"]["mrr"] >= 0.8
    assert not [f for f in extended_report["failures"] if f["id"] != "x14"] and set(m["by_type"]) == {"semantic", "compare", "contradiction", "hindi", "analytics", "parallel", "abstain"}
    assert all(t["correct"] == 1.0 for t in m["by_type"].values())


def test_agent_failures_and_forbidden_content_are_detected(tmp_path, monkeypatch):
    from app.agents import Assistant, nodes
    from app.storage import Scope

    s = Settings(data_dir=str(tmp_path / "d"), _env_file=None)
    k = Knowledge(s)
    p = tmp_path / "a.txt"
    p.write_text("Employees accrue 25 days of annual leave per year.")
    scope = Scope("t", "w")
    k.ingest_file(scope, p, "a.txt")
    monkeypatch.setattr(nodes.Agents, "semantic", lambda self, st: (_ for _ in ()).throw(RuntimeError("boom")))
    ds = EvalDataset(name="x", cases=[EvalCase(id="c1", question="How many days of annual leave do employees accrue?", expected_keywords=["25 days"], forbidden_keywords=["could not complete"], expected_sources=["a.txt"])])
    rep = run_eval(Assistant(k), scope, ANONYMOUS, ds)
    r = rep["cases"][0]
    assert r["agent_ok"] is False and r["forbidden_ok"] is False and r["sources_ok"] is False and rep["metrics"]["agents"]["execution_success_rate"] == 0.0
    assert rep["failures"][0]["why"] == ["keywords_ok", "forbidden_ok", "sources_ok", "agent execution failed"]


def test_failed_tool_calls_are_counted_but_declined_ones_are_not(tmp_path):
    from app.agents import Assistant
    from app.storage import Scope

    s = Settings(data_dir=str(tmp_path / "d"), _env_file=None)
    k = Knowledge(s)
    csv = tmp_path / "s.csv"
    csv.write_text("Name,Sales,Cost\na,1,2\nb,3,4\n")
    k.ingest_file(Scope("t", "w"), csv, "s.csv")
    a = Assistant(k)
    ds = EvalDataset(name="x", cases=[EvalCase(id="c1", question="What is the highest value?", type="data")])    # ambiguous column -> the agent asks, the tool "declines"
    rep = run_eval(a, Scope("t", "w"), ANONYMOUS, ds)
    ag = rep["metrics"]["agents"]
    assert ag["failed_tool_calls"] == 0 and ag["declined_tool_calls"] == 1 and ag["tool_calls"] == 1 and ag["failed_tool_call_rate"] == 0.0
    try:
        a.tools.call("exact_search", None, term="")                                                         # invalid arguments -> a real tool failure
    except Exception:
        pass
    reg = k.obs.metrics.render().decode()
    assert 'mka_tool_calls_total{outcome="declined",tool="table_query"} 1.0' in reg and 'mka_tool_calls_total{outcome="error",tool="exact_search"} 1.0' in reg


def test_eval_report_exports_pdf_docx_xlsx_with_sources_and_timestamps(extended_report):
    import io

    import docx
    import openpyxl
    import pymupdf

    from app.eval.report import to_doc
    from app.reports.render import render

    doc = to_doc(extended_report)
    assert doc.timestamp.startswith(extended_report["created"][:10]) and {s["name"] for s in doc.sources} >= {"leave_policy.txt", "orders.csv"}
    pdf = "\n".join(p.get_text() for p in pymupdf.open(stream=render(doc, "pdf"), filetype="pdf"))
    assert "Evaluation report" in pdf and "citation_accuracy" in pdf and "Failed tool calls" in pdf and "leave_policy.txt" in pdf and extended_report["created"][:10] in pdf
    d = docx.Document(io.BytesIO(render(doc, "docx")))
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert "Agent execution" in [p.text for p in d.paragraphs] and "Execution success rate" in cells and "hindi_policy.txt" in cells
    wb = openpyxl.load_workbook(io.BytesIO(render(doc, "xlsx")))
    assert wb.sheetnames == ["Report", "Sources"] and len(wb["Report"]._charts) == 1 and {r[0].value for r in wb["Sources"].iter_rows(min_row=2)} >= {"leave_policy.txt"}


def test_eval_api_extended_run_and_exports(tmp_path):
    s = Settings(data_dir=str(tmp_path / "d"), auth_enabled=True, jwt_secret="e" * 40, rate_limit_auth="1000/minute", rate_limit_ingest="1000/minute", _env_file=None)
    c = TestClient(create_app(s))
    adm = {"Authorization": "Bearer " + c.post("/auth/signup", json={"tenant_name": "Tenant", "email": "a@a.io", "password": "long enough pw"}).json()["access_token"]}
    r = c.post("/eval/run", json={"sample": True, "extended": True, "name": "ext"}, headers=adm)
    assert r.status_code == 200, r.text
    rid = r.json()["id"]
    assert r.json()["dataset"] == "platform-v3-extended" and r.json()["metrics"]["agents"]["execution_success_rate"] == 1.0
    for fmt, magic in (("pdf", b"%PDF-"), ("docx", b"PK"), ("xlsx", b"PK")):
        f = c.get(f"/eval/runs/{rid}/report.{fmt}", headers=adm)
        assert f.status_code == 200 and f.content[:len(magic)] == magic and fmt in f.headers["content-disposition"], fmt
    assert c.get(f"/eval/runs/{rid}/report.md", headers=adm).text.startswith("# Evaluation report")          # md/html routes still win over the generic format route
    assert "<!doctype html>" in c.get(f"/eval/runs/{rid}/report.html", headers=adm).text
    assert c.get(f"/eval/runs/{rid}/report.exe", headers=adm).status_code == 404
    other = {"Authorization": "Bearer " + c.post("/auth/signup", json={"tenant_name": "Other", "email": "b@b.io", "password": "long enough pw"}).json()["access_token"]}
    assert c.get(f"/eval/runs/{rid}/report.pdf", headers=other).status_code == 404
